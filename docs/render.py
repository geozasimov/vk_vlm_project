"""Мини-рендерер документа: один список блоков -> HTML (-> PDF через headless Chrome/Edge) и .docx.

Блоки: ("h1"|"h2"|"h3", текст) · ("p", текст) · ("ul", [пункты]) · ("ol", [пункты])
       ("table", {"head": [...], "rows": [[...]], "widths": [доли], "bold_rows": {i}}) · ("img", путь, подпись)
       ("note", текст) · ("pagebreak",)
В тексте поддерживается **жирный** и `код`.
"""
from __future__ import annotations

import html
import re
import shutil
import subprocess
from pathlib import Path

INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")

CSS = """
@page { size: A4; margin: 18mm 17mm 18mm 17mm; }
:root { --ink:#0b0b0b; --ink2:#52514e; --muted:#898781; --rule:#e1e0d9; --accent:#2a78d6; --wash:#f0efec; }
body { font-family: "Segoe UI", "Helvetica Neue", Arial, sans-serif; color: var(--ink); background:#fff;
       font-size: 10.5pt; line-height: 1.45; }
h1 { font-size: 20pt; margin: 0 0 4pt; line-height:1.2; }
.subtitle { color: var(--ink2); font-size: 11pt; margin: 0 0 14pt; }
h2 { font-size: 14pt; margin: 18pt 0 6pt; padding-bottom: 3pt; border-bottom: 1px solid var(--rule);
     break-after: avoid; }
h3 { font-size: 11.5pt; margin: 12pt 0 4pt; break-after: avoid; }
p { margin: 0 0 6pt; text-align: left; }
ul, ol { margin: 0 0 6pt 16pt; padding: 0; } li { margin: 0 0 2pt; }
code { font-family: Consolas, "Courier New", monospace; font-size: 9pt; background: var(--wash);
       padding: 0 3px; border-radius: 3px; }
table { border-collapse: collapse; width: 100%; margin: 4pt 0 10pt; font-size: 9.5pt; break-inside: avoid; }
th { text-align: left; color: var(--ink2); font-weight: 600; border-bottom: 1.5px solid var(--ink2);
     padding: 4pt 6pt; }
td { border-bottom: 1px solid var(--rule); padding: 3.5pt 6pt; vertical-align: top; }
b { font-weight: 700; }
tr.b td { font-weight: 700; background: #eef4fc; }
td.num, th.num { text-align: right; font-variant-numeric: tabular-nums; }
figure { margin: 6pt 0 12pt; break-inside: avoid; } figure img { width: 100%; }
figcaption { color: var(--ink2); font-size: 9pt; margin-top: 3pt; }
.note { border-left: 3px solid var(--accent); background: #f5f8fd; padding: 6pt 9pt; margin: 6pt 0 10pt; }
.pb { break-after: page; }
"""

NUM = re.compile(r"^[−+\-]?\d[\d\s.,]*%?$|^—$")


def _inline_html(t: str) -> str:
    out = []
    for part in INLINE.split(t):
        if part.startswith("**") and part.endswith("**"):
            out.append(f"<b>{html.escape(part[2:-2])}</b>")
        elif part.startswith("`") and part.endswith("`"):
            out.append(f"<code>{html.escape(part[1:-1])}</code>")
        else:
            out.append(html.escape(part))
    return "".join(out)


def to_html(title: str, subtitle: str, blocks: list, base: Path) -> str:
    body = [f"<h1>{_inline_html(title)}</h1>", f"<p class='subtitle'>{_inline_html(subtitle)}</p>"]
    for b in blocks:
        kind = b[0]
        if kind in ("h1", "h2", "h3"):
            body.append(f"<{kind}>{_inline_html(b[1])}</{kind}>")
        elif kind == "p":
            body.append(f"<p>{_inline_html(b[1])}</p>")
        elif kind == "note":
            body.append(f"<div class='note'>{_inline_html(b[1])}</div>")
        elif kind in ("ul", "ol"):
            body.append(f"<{kind}>" + "".join(f"<li>{_inline_html(x)}</li>" for x in b[1]) + f"</{kind}>")
        elif kind == "table":
            t = b[1]
            widths = t.get("widths")
            cols = "".join(f"<col style='width:{w * 100:.0f}%'>" for w in widths) if widths else ""
            head = "".join(f"<th class='{'num' if i and t.get('num_cols', True) else ''}'>{_inline_html(h)}</th>"
                           for i, h in enumerate(t["head"]))
            rows = []
            for ri, r in enumerate(t["rows"]):
                cls = " class='b'" if ri in t.get("bold_rows", set()) else ""
                cells = "".join(
                    f"<td class='{'num' if NUM.match(str(c).strip('* ')) and ci else ''}'>{_inline_html(str(c))}</td>"
                    for ci, c in enumerate(r))
                rows.append(f"<tr{cls}>{cells}</tr>")
            body.append(f"<table><colgroup>{cols}</colgroup><thead><tr>{head}</tr></thead>"
                        f"<tbody>{''.join(rows)}</tbody></table>")
        elif kind == "img":
            if not Path(b[1]).exists():
                continue
            rel = Path(b[1]).resolve().as_uri()
            body.append(f"<figure><img src='{rel}'><figcaption>{_inline_html(b[2])}</figcaption></figure>")
        elif kind == "pagebreak":
            body.append("<div class='pb'></div>")
    return (f"<!doctype html><html lang='ru'><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
            f"<style>{CSS}</style></head><body>{''.join(body)}</body></html>")


def find_browser() -> str | None:
    for c in [r"C:\Program Files\Google\Chrome\Application\chrome.exe",
              r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe", shutil.which("chrome"),
              shutil.which("msedge")]:
        if c and Path(c).exists():
            return c
    return None


def html_to_pdf(html_path: Path, pdf_path: Path):
    browser = find_browser()
    if not browser:
        print("Chrome/Edge не найден — PDF не собран")
        return
    subprocess.run([browser, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={pdf_path}", html_path.resolve().as_uri()], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)


def to_docx(title: str, subtitle: str, blocks: list, path: Path):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10.5)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11.5)):
        s = doc.styles[name]
        s.font.name, s.font.size, s.font.color.rgb = "Calibri", Pt(size), RGBColor(0x0B, 0x0B, 0x0B)

    def add_inline(par, text, bold=False):
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = par.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("`") and part.endswith("`"):
                r = par.add_run(part[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(9)
            else:
                r = par.add_run(part)
            if bold:
                r.bold = True

    def shade(cell, hex_fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tcPr.append(shd)

    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph()
    add_inline(p, subtitle)
    width_cm = 21 - 3.6
    for b in blocks:
        kind = b[0]
        if kind == "h1":
            doc.add_heading(b[1].replace("**", ""), level=1)
        elif kind == "h2":
            doc.add_heading(b[1].replace("**", ""), level=1)
        elif kind == "h3":
            doc.add_heading(b[1].replace("**", ""), level=2)
        elif kind in ("p", "note"):
            add_inline(doc.add_paragraph(), b[1])
        elif kind in ("ul", "ol"):
            for x in b[1]:
                add_inline(doc.add_paragraph(style="List Bullet" if kind == "ul" else "List Number"), x)
        elif kind == "table":
            t = b[1]
            tbl = doc.add_table(rows=1 + len(t["rows"]), cols=len(t["head"]))
            tbl.style = "Table Grid"
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            widths = t.get("widths") or [1 / len(t["head"])] * len(t["head"])
            for ri, row in enumerate([t["head"]] + t["rows"]):
                for ci, val in enumerate(row):
                    cell = tbl.cell(ri, ci)
                    cell.width = Cm(width_cm * widths[ci])
                    cell.paragraphs[0].text = ""
                    add_inline(cell.paragraphs[0], str(val), bold=ri == 0 or (ri - 1) in t.get("bold_rows", set()))
                    for r in cell.paragraphs[0].runs:
                        r.font.size = Pt(9)
                    if ri == 0:
                        shade(cell, "F0EFEC")
                    elif (ri - 1) in t.get("bold_rows", set()):
                        shade(cell, "EEF4FC")
            doc.add_paragraph()
        elif kind == "img":
            if not Path(b[1]).exists():
                continue
            doc.add_picture(str(b[1]), width=Cm(width_cm))
            cap = doc.add_paragraph()
            r = cap.add_run(b[2].replace("**", "").replace("`", ""))
            r.italic, r.font.size = True, Pt(9)
        elif kind == "pagebreak":
            doc.add_page_break()
    doc.save(path)
